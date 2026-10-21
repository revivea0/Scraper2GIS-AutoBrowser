import asyncio
import json
import os
import re
import subprocess
import tempfile
import time
import urllib.parse
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from playwright.async_api import Error as PlaywrightError, Page, Route, async_playwright
from playwright._impl._driver import compute_driver_executable, get_driver_env


COLUMNS = ('ID', 'Название', 'Адрес', 'Рубрика', 'Телефоны', 'Сайт', 'Telegram', 'ВКонтакте', 'Ссылка 2GIS')


class ParsingStopped(Exception):
    pass


class TwoGisParser:
    def __init__(self, log_callback, stats_callback):
        self.log = log_callback
        self.stats = stats_callback
        self.results = {}
        self.processed_firm_ids = set()
        self.is_running = True
        self.current_page = 1
        self._pending_responses = set()
        self._firm_events = {}
        self._page_ids = {}
        self._search_context = None
        self._page_complete = False
        self._backup_path = None
        self._blocked_reason = ''

    async def _sleep(self, duration):
        deadline = time.monotonic() + duration
        while self.is_running and time.monotonic() < deadline:
            await asyncio.sleep(min(0.1, max(0, deadline - time.monotonic())))

    async def _wait(self, awaitable, timeout):
        task = asyncio.ensure_future(awaitable)
        deadline = time.monotonic() + timeout
        try:
            while self.is_running:
                if self._blocked_reason:
                    raise RuntimeError(self._blocked_reason)
                remaining = deadline - time.monotonic()
                if remaining <= 0:
                    raise TimeoutError('Превышено время ожидания 2GIS')
                done, _ = await asyncio.wait({task}, timeout=min(0.1, remaining))
                if done:
                    return task.result()
            raise ParsingStopped()
        finally:
            if not task.done():
                task.cancel()
                await asyncio.gather(task, return_exceptions=True)

    def _generate_url(self, city, query, scope, detail):
        if scope == 'district' and detail:
            search = f'{city} район {detail} {query}'
        elif scope == 'street' and detail:
            search = f'{city} улица {detail} {query}'
        else:
            search = f'{city} {query}'
        return 'https://2gis.ru/search/' + urllib.parse.quote(search, safe='')

    def get_blocked_patterns(self):
        return ['/_/log', '/_/metrics', 'google-analytics.com', 'counter.yadro.ru',
                'tns-counter.ru', 'mc.yandex.ru', 'catalog.api.2gis.ru/3.0/ads/',
                'tile', 'bss.2gis.', 'styles.api.2gis.', 'video-pr.api.2gis.',
                'api.photo.2gis.', 'market-backend.api.2gis.']

    async def _route_intercept(self, route: Route):
        if route.request.resource_type in ('image', 'media', 'font') or any(
            pattern in route.request.url for pattern in self.get_blocked_patterns()
        ):
            await route.abort()
        else:
            await route.continue_()

    async def _kill_popups(self, page: Page):
        for text in ('Хорошо', 'Да, всё верно', 'Понятно', 'Принять', 'Согласен'):
            button = page.get_by_role('button', name=text, exact=True).first
            try:
                if await button.is_visible():
                    await self._wait(button.click(timeout=1000), 2)
            except PlaywrightError:
                continue

    @staticmethod
    def _card_ids(hrefs):
        ids = []
        for href in hrefs:
            parsed = urllib.parse.urlsplit(href or '')
            if parsed.hostname not in (None, '2gis.ru', 'www.2gis.ru'):
                continue
            match = re.search(r'/firm/(\d+)(?:[_/]|$)', parsed.path)
            if match and match.group(1) not in ids:
                ids.append(match.group(1))
        return ids

    @staticmethod
    def _next_page_url(hrefs, base_url, page_num):
        base = urllib.parse.urlsplit(base_url)
        expected = base.path.rstrip('/') + f'/page/{page_num}'
        for href in hrefs:
            candidate = urllib.parse.urljoin(base_url, href)
            parsed = urllib.parse.urlsplit(candidate)
            if parsed.netloc == base.netloc and parsed.path.rstrip('/') == expected:
                return candidate
        return None

    def _load_backup(self, backup_path, search_context=None):
        self.log(f'Чтение файла: {Path(backup_path).name}')
        loaded = {}
        wb = openpyxl.load_workbook(backup_path, read_only=True, data_only=False)
        try:
            metadata = None
            if '_scraper' in wb.sheetnames:
                metadata = json.loads(wb['_scraper']['A1'].value)
                if (not isinstance(metadata, dict) or metadata.get('version') != 1
                        or type(metadata.get('page')) is not int or metadata['page'] < 1
                        or type(metadata.get('page_complete')) is not bool):
                    raise ValueError('Неизвестный или повреждённый формат бэкапа')
                if search_context and metadata.get('search') and metadata['search'] != search_context:
                    raise ValueError('Бэкап создан с другими параметрами поиска: выберите исходный город, рубрику и область')
            ws = wb.worksheets[0]
            rows = iter(ws.iter_rows())
            header = next(rows, ())
            headers = [cell.value for cell in header]
            if 'ID' not in headers:
                raise ValueError('В бэкапе отсутствует столбец ID')
            for cells in rows:
                row = {key: cell.value for key, cell in zip(headers, cells) if key in COLUMNS}
                raw_id = row.get('ID')
                if raw_id is None or raw_id == '':
                    continue
                if isinstance(raw_id, (int, float)) and len(str(int(raw_id))) > 15:
                    raise ValueError('Длинный ID сохранён числом и мог быть округлён Excel. Нужен исходный текстовый ID')
                org_id = str(raw_id).strip().split('_')[0]
                if not org_id.isascii() or not org_id.isdigit():
                    raise ValueError(f'Некорректный ID в строке {cells[0].row}')
                record = loaded.setdefault(org_id, {'ID': org_id})
                self._merge_record(record, {key: str(value) for key, value in row.items() if key != 'ID' and value is not None})
                record['Ссылка 2GIS'] = f'https://2gis.ru/firm/{org_id}'
            if not loaded:
                raise ValueError('В бэкапе нет записей организаций')
            if metadata:
                page = metadata['page'] + int(bool(metadata.get('page_complete')))
            else:
                match = re.match(r'(BACKUP|STOPPED)_page_(\d+)_', Path(backup_path).name)
                page = max(1, int(match.group(2))) if match else 1
                if match and match.group(1) == 'BACKUP':
                    page += 1
                self.log('Старый файл без параметров поиска. Убедитесь, что город и рубрика совпадают с исходными.')
        finally:
            wb.close()
        self.results.update(loaded)
        self.processed_firm_ids.update(loaded)
        self._backup_path = Path(backup_path).resolve()
        self.log(f'Восстановлено {len(loaded)} записей. Продолжаем со страницы {page}.')
        return page

    async def _selected_page(self, page):
        return await self._wait(page.evaluate(r'''() => {
            const current = document.querySelector('[aria-current="page"]');
            if (current && /^\d+$/.test(current.textContent.trim())) return Number(current.textContent.trim());
            for (const link of document.querySelectorAll('a[href*="/page/"]')) {
                const selected = Array.from(link.parentElement.children).find(node =>
                    node.tagName !== 'A' && /^\d+$/.test(node.textContent.trim()));
                if (selected) return Number(selected.textContent.trim());
            }
            return null;
        }'''), 3)

    async def _go_to_page(self, page, base_url, target, previous):
        while self.is_running:
            selected = await self._selected_page(page)
            if selected == target:
                return
            hrefs = await self._wait(page.locator('a[href*="/page/"]').evaluate_all(
                '(nodes) => nodes.map(node => node.getAttribute("href"))'), 3)
            candidates = []
            for href in hrefs:
                match = re.search(r'/page/(\d+)(?:[/?#]|$)', href or '')
                if match and (selected or 1) < int(match.group(1)) <= target:
                    candidates.append(int(match.group(1)))
            if not candidates:
                raise RuntimeError(f'Невозможно перейти на страницу {target}: проверьте, что запрос совпадает с бэкапом.')
            step = max(candidates)
            next_url = self._next_page_url(hrefs, base_url, step)
            if not next_url:
                raise RuntimeError('Не найден адрес следующей страницы 2GIS')
            def is_search_response(response):
                url = urllib.parse.urlsplit(response.url)
                return (url.hostname == 'catalog.api.2gis.ru' and url.path == '/3.0/items'
                        and urllib.parse.parse_qs(url.query).get('page') == [str(step)])

            async with page.expect_response(is_search_response, timeout=20000) as search:
                await self._wait(page.locator('a').evaluate_all('''(nodes, target) => {
                    const link = nodes.find(node => node.href === target);
                    if (!link) throw new Error('Не найдена ссылка пагинации');
                    link.click();
                }''', next_url), 5)
                response = await self._wait(search.value, 22)
            if response.status != 200:
                raise RuntimeError(f'Не удалось загрузить страницу {step}: HTTP {response.status}. Повторите позже.')
            data = await self._wait(response.json(), 10)
            result = data.get('result')
            if not isinstance(result, dict) or not isinstance(result.get('items'), list):
                raise RuntimeError('2GIS не вернул список организаций. Прогресс сохранится.')
            ids = [str(item.get('id', '')).split('_')[0] for item in result['items']]
            self._page_ids[step] = list(dict.fromkeys(org_id for org_id in ids if org_id.isascii() and org_id.isdigit()))
            if result['items'] and not self._page_ids[step]:
                raise RuntimeError('Неизвестный формат списка организаций 2GIS')
            previous = await self._page_cards(page, previous, expected_page=step)
            if step == target:
                return
        raise ParsingStopped()

    async def _page_cards(self, page, previous=(), expected_page=None):
        expected_page = expected_page or self.current_page
        deadline = time.monotonic() + 20
        while self.is_running and time.monotonic() < deadline:
            if self._blocked_reason:
                raise RuntimeError(self._blocked_reason)
            hrefs = await self._wait(page.locator('a[href*="/firm/"]').evaluate_all(
                '(nodes) => nodes.map(node => node.getAttribute("href"))'), 3)
            ids = self._card_ids(hrefs)
            selected = await self._selected_page(page)
            expected_ids = self._page_ids.get(expected_page)
            if expected_ids is not None:
                if not expected_ids:
                    return []
                if set(expected_ids) <= set(ids) and selected == expected_page:
                    return expected_ids
            elif ids and set(ids) != set(previous) and (selected == expected_page or (selected is None and expected_page == 1)):
                self._page_ids[expected_page] = ids
                return ids
            text = await self._wait(page.locator('body').inner_text(timeout=2000), 3)
            if not expected_ids and re.search(r'ничего не найдено|ничего не нашли|нет результатов|не нашлось', text, re.I):
                return []
            if re.search(r'подтвердите.{0,60}человек|проверка.{0,30}безопасности|captcha', text, re.I):
                raise RuntimeError('2GIS запросил проверку доступа. Сбор остановлен, прогресс сохранится.')
            await self._sleep(0.1)
        if not self.is_running:
            raise ParsingStopped()
        raise RuntimeError('Карточки 2GIS не загрузились или страница не сменилась. Прогресс сохранится; повторите позже.')

    async def _collect_firm(self, page, firm_id):
        event = self._firm_events.setdefault(firm_id, asyncio.Event())
        try:
            for attempt in range(2):
                if not self.is_running:
                    raise ParsingStopped()
                if self._blocked_reason:
                    raise RuntimeError(self._blocked_reason)
                try:
                    if attempt:
                        self.log(f'Повтор загрузки организации {firm_id}...')
                        await self._sleep(0.5)
                        url = urllib.parse.urlsplit(page.url)
                        path = url.path.split('/firm/')[0] + '/firm/' + firm_id
                        retry_url = urllib.parse.urlunsplit((url.scheme, url.netloc, path, '', ''))
                        await self._wait(page.goto(retry_url, wait_until='domcontentloaded', timeout=20000), 22)
                    else:
                        await self._wait(page.locator('a[href*="/firm/"]').evaluate_all(r'''(nodes, id) => {
                            const card = nodes.find(node => {
                                const url = new URL(node.href);
                                const match = url.pathname.match(/\/firm\/(\d+)(?:[_/]|$)/);
                                return url.hostname === '2gis.ru' && match && match[1] === id;
                            });
                            if (!card) throw new Error('Карточка исчезла из списка');
                            card.click();
                        }''', firm_id), 6)
                    await self._wait(event.wait(), 10)
                    return firm_id in self.processed_firm_ids
                except (TimeoutError, PlaywrightError) as error:
                    if page.is_closed():
                        raise RuntimeError('Окно браузера закрыто. Прогресс сохранится.') from error
            return False
        finally:
            self._firm_events.pop(firm_id, None)

    async def _launch_browser(self, playwright):
        for channel in ('msedge', 'chrome'):
            try:
                return await self._wait(playwright.chromium.launch(headless=False, channel=channel), 20)
            except PlaywrightError:
                continue
        if not Path(playwright.chromium.executable_path).is_file():
            self.log('Браузер не найден. Скачиваем Chromium; дождитесь завершения установки...')
            # The bundled Node CLI also works inside a PyInstaller executable.
            node, cli = compute_driver_executable()
            process = await asyncio.create_subprocess_exec(
                node, cli, 'install', 'chromium', env=get_driver_env(),
                stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.STDOUT,
                **({'creationflags': subprocess.CREATE_NO_WINDOW} if os.name == 'nt' else {}))
            try:
                output, _ = await self._wait(process.communicate(), 600)
                if process.returncode:
                    raise RuntimeError('Не удалось скачать Chromium. Проверьте интернет и повторите сбор. '
                                       + output.decode('utf-8', errors='replace')[-1500:])
            finally:
                if process.returncode is None:
                    process.kill()
                    await process.communicate()
            self.log('Chromium установлен. Запускаем браузер...')
        return await self._wait(playwright.chromium.launch(headless=False), 20)

    async def run(self, city, query, scope, detail, save_dir, backup_path):
        if not self.is_running:
            return ''
        if not all(isinstance(value, str) for value in (city, query, scope, detail, save_dir, backup_path)):
            raise ValueError('Параметры поиска должны быть строками')
        city, query, detail = city.strip(), query.strip(), detail.strip()
        if not city or not query or scope not in ('city', 'district', 'street') or (scope != 'city' and not detail):
            raise ValueError('Укажите город, рубрику и название выбранного района или улицы')
        directory = Path(save_dir or '.').resolve()
        directory.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryFile(dir=directory):
            pass
        self.results.clear()
        self.processed_firm_ids.clear()
        self._page_ids.clear()
        self._blocked_reason = ''
        self._search_context = {'city': city, 'query': query, 'scope': scope, 'detail': detail if scope != 'city' else ''}
        self.current_page = self._load_backup(backup_path, self._search_context) if backup_path else 1
        self._page_complete = False
        self.stats(len(self.results), self.current_page)
        finished = False
        filepath = ''
        browser = None
        page = None
        context = None
        async with async_playwright() as playwright:
            try:
                browser = await self._launch_browser(playwright)
                context = await browser.new_context(viewport={'width': 1280, 'height': 800}, service_workers='block')
                page = await context.new_page()
                page.set_default_timeout(5000)
                await page.route('**/*', self._route_intercept)
                page.on('response', self._schedule_api_response)
                self.log('Открываем карту...')
                await self._wait(page.goto(self._generate_url(city, query, scope, detail), wait_until='domcontentloaded', timeout=20000), 22)
                base_url = page.url.split('?')[0].split('/firm/')[0].split('/page/')[0].rstrip('/')
                if self.current_page > 1:
                    initial_ids = await self._page_cards(page, expected_page=1)
                    await self._go_to_page(page, base_url, self.current_page, initial_ids)
                previous = ()
                while self.is_running:
                    self._page_complete = False
                    self.stats(len(self.results), self.current_page)
                    self.log(f'Сканируем страницу {self.current_page}...')
                    ids = await self._page_cards(page, previous)
                    if not ids:
                        self.log('В выдаче нет организаций.')
                        finished = True
                        self._page_complete = True
                        break
                    await self._kill_popups(page)
                    pending = [org_id for org_id in ids if org_id not in self.processed_firm_ids]
                    self.log(f'Сбор контактов: {len(pending)} новых организаций.')
                    for org_id in pending:
                        if not await self._collect_firm(page, org_id):
                            raise RuntimeError(f'Не удалось получить организацию {org_id}. Страница {self.current_page} будет повторена при восстановлении.')
                    await self._wait_for_responses()
                    self._page_complete = True
                    self.stats(len(self.results), self.current_page)
                    checkpoint = self._export_to_excel(city, query, str(directory), is_backup=True, current_page=self.current_page)
                    if self.results and not checkpoint:
                        raise RuntimeError('Не удалось записать бэкап. Сбор остановлен, чтобы не потерять прогресс.')
                    if not self.is_running:
                        break
                    hrefs = await self._wait(page.locator('a[href*="/page/"]').evaluate_all(
                        '(nodes) => nodes.map(node => node.getAttribute("href"))'), 3)
                    next_url = self._next_page_url(hrefs, base_url, self.current_page + 1)
                    if not next_url:
                        finished = True
                        self.log('Конец выдачи 2GIS. Парсинг успешно завершен.')
                        break
                    previous = tuple(ids)
                    self.current_page += 1
                    self._page_complete = False
                    await self._go_to_page(page, base_url, self.current_page, ids)
            except ParsingStopped:
                self.log('Сбор остановлен. Сохраняем полученные данные.')
            except (PlaywrightError, TimeoutError, RuntimeError, OSError) as error:
                self.log(f'Ошибка: {error}')
            finally:
                if page is not None:
                    page.remove_listener('response', self._schedule_api_response)
                await self._wait_for_responses()
                self.stats(len(self.results), self.current_page)
                if self.results:
                    self.log('Сохранение итоговой таблицы...')
                    filepath = self._export_to_excel(city, query, str(directory), current_page=self.current_page,
                                                     prefix_override='' if finished else 'STOPPED_')
                if browser is not None:
                    try:
                        if page is not None and not page.is_closed():
                            await asyncio.wait_for(page.unroute_all(behavior='ignoreErrors'), timeout=3)
                            await asyncio.wait_for(page.close(run_before_unload=False), timeout=3)
                    except (PlaywrightError, TimeoutError):
                        pass
                    if context is not None:
                        try:
                            await asyncio.wait_for(context.close(), timeout=3)
                        except (PlaywrightError, TimeoutError):
                            pass
                    try:
                        await asyncio.wait_for(browser.close(), timeout=10)
                    except (PlaywrightError, TimeoutError):
                        self.log('Браузер не ответил при закрытии. Данные уже сохранены или доступны для повторного сохранения.')
        return filepath

    def _schedule_api_response(self, response):
        parsed = urllib.parse.urlsplit(response.url)
        if parsed.hostname != 'catalog.api.2gis.ru' or parsed.path != '/3.0/items/byid':
            return
        if response.status in (403, 429):
            self._blocked_reason = f'2GIS ограничил доступ (HTTP {response.status}). Повторите сбор позже.'
            return
        if response.status == 200:
            task = asyncio.create_task(self._intercept_api_response(response))
            self._pending_responses.add(task)
            task.add_done_callback(self._pending_responses.discard)

    async def _wait_for_responses(self):
        if self._pending_responses:
            _, pending = await asyncio.wait(self._pending_responses, timeout=10)
            for task in pending:
                task.cancel()
            if pending:
                await asyncio.gather(*pending, return_exceptions=True)
                self.log('Ошибка: часть ответов 2GIS не загрузилась за 10 секунд. Сохраняем полученные данные.')

    async def _intercept_api_response(self, response):
        try:
            data = await response.json()
            items = (data.get('result') or {}).get('items') or []
            for item in items:
                try:
                    org_id = str(item.get('id') or '').split('_')[0]
                    if not org_id.isascii() or not org_id.isdigit():
                        continue
                    record = {'ID': org_id, 'Название': item.get('name') or 'Без названия',
                              'Адрес': item.get('address_name') or 'Без адреса',
                              'Рубрика': self._extract_category(item), 'Телефоны': self._extract_phone(item),
                              'Сайт': self._extract_specific_contact(item, 'website'),
                              'Telegram': self._extract_specific_contact(item, 'telegram'),
                              'ВКонтакте': self._extract_specific_contact(item, 'vkontakte'),
                              'Ссылка 2GIS': f'https://2gis.ru/firm/{org_id}'}
                    is_new = org_id not in self.results
                    saved = self.results.setdefault(org_id, {})
                    self._merge_record(saved, record)
                    self.processed_firm_ids.add(org_id)
                    if org_id in self._firm_events:
                        self._firm_events[org_id].set()
                    if is_new:
                        self.log(f'+ {record["Название"]}')
                        self.stats(len(self.results), -1)
                except (AttributeError, TypeError, ValueError) as error:
                    self.log(f'Ошибка обработки карточки: {error}')
        except (PlaywrightError, AttributeError, TypeError, ValueError) as error:
            self.log(f'Ошибка чтения ответа 2GIS: {error}')

    @staticmethod
    def _merge_record(saved, incoming):
        for key, value in incoming.items():
            if not value or (key in saved and value in ('Нет телефона', 'Без адреса')):
                continue
            if key in ('Телефоны', 'Сайт', 'Telegram', 'ВКонтакте') and value != 'Нет телефона':
                previous = saved.get(key, '')
                values = [] if previous == 'Нет телефона' else previous.split(', ')
                saved[key] = ', '.join(dict.fromkeys(part for part in values + str(value).split(', ') if part))
            else:
                saved[key] = value

    def _contacts(self, item, contact_type):
        values = []
        for group in item.get('contact_groups') or []:
            for contact in group.get('contacts') or []:
                if contact.get('type') == contact_type:
                    value = contact.get('value') or contact.get('url')
                    if value and str(value) not in values:
                        values.append(str(value))
        return ', '.join(values)

    def _extract_category(self, item):
        return ', '.join(dict.fromkeys(rubric['name'] for rubric in item.get('rubrics') or [] if rubric.get('name')))

    def _extract_phone(self, item):
        return self._contacts(item, 'phone') or 'Нет телефона'

    def _extract_specific_contact(self, item, contact_type):
        return self._contacts(item, contact_type)

    def _export_to_excel(self, city, query, save_dir, is_backup=False, current_page=1, prefix_override=''):
        if not self.results:
            return ''
        temp_path = None
        wb = None
        try:
            suffix = f'{city}_{query}'.replace(' ', '_')
            if is_backup:
                filename = f'BACKUP_{suffix}.xlsx'
            elif prefix_override:
                filename = f'{prefix_override}page_{current_page}_{suffix}.xlsx'
            else:
                filename = f'Parser_{suffix}.xlsx'
            filename = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', filename)
            filename = filename[:-5].encode('utf-8')[:180].decode('utf-8', errors='ignore') + '.xlsx'
            filepath = Path(save_dir or '.').resolve() / filename
            if filepath == self._backup_path:
                filepath = filepath.with_name(filepath.stem + '_resumed.xlsx')
            if not is_backup:
                base = filepath
                index = 2
                while filepath.exists():
                    filepath = base.with_name(f'{base.stem}_{index}.xlsx')
                    index += 1
            with tempfile.NamedTemporaryFile(dir=filepath.parent, prefix='.scraper-', suffix='.xlsx', delete=False) as temp:
                temp_path = temp.name
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'Организации'
            ws.append(COLUMNS)
            fill = PatternFill('solid', fgColor='335CDB')
            for cell in ws[1]:
                cell.fill = fill
                cell.font = Font(color='FFFFFF', bold=True)
                cell.alignment = Alignment(vertical='center')
            widths = [len(title) for title in COLUMNS]
            align = Alignment(vertical='top', wrap_text=True)
            for row_index, record in enumerate(self.results.values(), 2):
                values = [re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', str(record.get(key) or ''))[:32767] for key in COLUMNS]
                ws.append(values)
                for index, value in enumerate(values):
                    cell = ws.cell(row=row_index, column=index + 1)
                    cell.data_type = 's'
                    cell.alignment = align
                    widths[index] = max(widths[index], len(values[index]))
            for index, width in enumerate(widths, 1):
                ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = min(max(width + 2, 16), 52)
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            metadata = wb.create_sheet('_scraper')
            metadata['A1'] = json.dumps({'version': 1, 'search': self._search_context,
                                         'page': current_page, 'page_complete': self._page_complete}, ensure_ascii=False)
            metadata.sheet_state = 'hidden'
            wb.save(temp_path)
            wb.close()
            os.replace(temp_path, filepath)
            temp_path = None
            if not is_backup:
                self.log(f'Таблица готова: {filepath.name}')
            return str(filepath)
        except (OSError, ValueError, TypeError) as error:
            self.log(f'Ошибка сохранения: {error}. Проверьте папку, свободное место и доступ к файлу.')
            return ''
        finally:
            if wb is not None:
                wb.close()
            if temp_path:
                try:
                    Path(temp_path).unlink(missing_ok=True)
                except OSError:
                    pass
