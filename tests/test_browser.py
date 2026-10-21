"""Opt in with SCRAPER_BROWSER_TESTS=1; all browser traffic is fulfilled locally."""
import asyncio
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from urllib.parse import parse_qs, urlsplit

import openpyxl
from playwright.async_api import BrowserType

from core.parser import TwoGisParser


@unittest.skipUnless(os.getenv('SCRAPER_BROWSER_TESTS') == '1', 'Set SCRAPER_BROWSER_TESTS=1 and install Chromium')
class BrowserTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.logs = []
        self.requests = []
        self.mode = 'success'
        self.failed_once = False
        self.stop_after = 0
        owner = self
        original_launch = BrowserType.launch

        async def launch(browser_type, **kwargs):
            kwargs.pop('channel', None)
            kwargs['headless'] = True
            return await original_launch(browser_type, **kwargs)

        self.launch_patch = patch.object(BrowserType, 'launch', launch)
        self.launch_patch.start()
        self.addCleanup(self.launch_patch.stop)

        class FixtureParser(TwoGisParser):
            def _generate_url(self, *args):
                return 'https://2gis.ru/search/fixture'

            async def _wait(self, awaitable, timeout):
                return await super()._wait(awaitable, 0.3 if timeout == 10 else timeout)

            async def _route_intercept(self, route):
                url = urlsplit(route.request.url)
                if url.hostname == '2gis.ru':
                    if owner.mode == 'empty':
                        html = '<html lang="ru"><body>Ничего не найдено</body></html>'
                    elif owner.mode == 'captcha':
                        html = '<html lang="ru"><body>Подтвердите, что вы человек</body></html>'
                    else:
                        html = Path(__file__).with_name('fixtures').joinpath('catalog.html').read_text()
                        if owner.mode == 'retry':
                            html = html.replace('<script>', '<script>window.delayFailedUrl=true;')
                        if owner.mode == 'transient_empty':
                            html = html.replace('<script>', '<script>window.fixtureEmpty=true; window.fixtureDelay=150;')
                        if owner.mode == 'delayed':
                            html = html.replace('<script>', '<script>window.fixtureDelay=150;')
                    await route.fulfill(status=200, content_type='text/html; charset=utf-8', body=html)
                elif url.hostname == 'catalog.api.2gis.ru':
                    query = parse_qs(url.query)
                    if url.path == '/3.0/items':
                        if owner.mode == 'search_limit':
                            await route.fulfill(status=429, body='{}')
                            return
                        number = int(query['page'][0])
                        ids = {1: ['101','102'], 2: ['201','202'], 3: ['301']}[number]
                        await route.fulfill(status=200, content_type='application/json', body=json.dumps(
                            {'result': {'items': [{'id': org_id} for org_id in ids]}}))
                        return
                    org_id = query['id'][0].split('_')[0]
                    owner.requests.append(org_id)
                    await asyncio.sleep(0.02)
                    if owner.mode == 'blocked':
                        await route.fulfill(status=429, body='{}')
                    elif org_id == '102' and (owner.mode == 'failure' or (owner.mode == 'retry' and not owner.failed_once)):
                        owner.failed_once = True
                        await route.fulfill(status=500, body='{}')
                    else:
                        payload = {'result': {'items': [{'id': org_id + '_session', 'name': 'Кафе ' + org_id,
                            'address_name': 'Тестовый адрес', 'rubrics': [{'name': 'Кафе'}],
                            'contact_groups': [{'contacts': [{'type': 'phone', 'value': '+79990000000'}]}]}]}}
                        await route.fulfill(status=200, content_type='application/json', body=json.dumps(payload))
                else:
                    await route.abort()

        def stats(count, page):
            if self.stop_after and count >= self.stop_after:
                self.parser.is_running = False

        self.parser = FixtureParser(self.logs.append, stats)

    async def run_parser(self, backup=''):
        return await self.parser.run('Тест', 'Кафе', 'city', '', self.tmp.name, backup)

    def checkpoint(self):
        path = Path(self.tmp.name, 'BACKUP_Тест_Кафе.xlsx')
        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            return json.loads(wb['_scraper']['A1'].value)
        finally:
            wb.close()

    async def test_three_pages_export_and_checkpoint(self):
        path = await self.run_parser()
        self.assertEqual(set(self.parser.results), {'101', '102', '201', '202', '301'})
        self.assertEqual(self.parser.processed_firm_ids, set(self.parser.results))
        self.assertTrue(Path(path).name.startswith('Parser_'))
        self.assertEqual(self.checkpoint()['page'], 3)
        self.assertTrue(self.checkpoint()['page_complete'])
        self.assertNotIn('999', self.requests)
        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            self.assertEqual(wb.active.max_row, 6)
        finally:
            wb.close()

    async def test_delayed_pagination_does_not_reuse_old_cards(self):
        self.mode = 'delayed'
        await self.run_parser()
        self.assertEqual(set(self.parser.results), {'101', '102', '201', '202', '301'})
        self.assertEqual(self.checkpoint()['page'], 3)

    async def test_transient_empty_state_does_not_override_verified_ids(self):
        self.mode = 'transient_empty'
        await self.run_parser()
        self.assertEqual(set(self.parser.results), {'101', '102', '201', '202', '301'})

    async def test_transient_failure_retries_without_losing_card(self):
        self.mode = 'retry'
        await self.run_parser()
        self.assertEqual(set(self.parser.results), {'101', '102', '201', '202', '301'})
        self.assertEqual(self.requests.count('102'), 2)

    async def test_failed_card_stops_and_resume_retries_same_page(self):
        self.mode = 'failure'
        path = await self.run_parser()
        self.assertEqual(set(self.parser.results), {'101'})
        self.assertNotIn('102', self.parser.processed_firm_ids)
        self.assertNotIn('201', self.requests)
        self.assertTrue(Path(path).name.startswith('STOPPED_page_1_'))
        self.mode = 'success'
        self.parser = type(self.parser)(self.logs.append, lambda *_: None)
        await self.run_parser(path)
        self.assertEqual(set(self.parser.results), {'101', '102', '201', '202', '301'})

    async def test_manual_stop_exports_received_rows(self):
        self.stop_after = 1
        path = await self.run_parser()
        self.assertEqual(set(self.parser.results), {'101'})
        self.assertTrue(Path(path).exists())
        self.assertNotIn('102', self.requests)

    async def test_empty_search_finishes_without_fake_file(self):
        self.mode = 'empty'
        self.assertEqual(await self.run_parser(), '')
        self.assertEqual(self.parser.results, {})
        self.assertTrue(any('нет организаций' in text for text in self.logs))

    async def test_access_limit_is_reported_and_not_retried(self):
        self.mode = 'blocked'
        await self.run_parser()
        self.assertTrue(any('429' in text for text in self.logs))
        self.assertEqual(self.requests, ['101'])

    async def test_search_limit_keeps_previous_page_and_resume_position(self):
        self.mode = 'search_limit'
        path = await self.run_parser()
        self.assertEqual(set(self.parser.results), {'101', '102'})
        self.assertNotIn('201', self.requests)
        self.assertTrue(Path(path).name.startswith('STOPPED_page_2_'))
        self.assertTrue(any('429' in text for text in self.logs))
        restored = TwoGisParser(self.logs.append, lambda *_: None)
        self.assertEqual(restored._load_backup(path), 2)

    async def test_captcha_stops_without_bypass(self):
        self.mode = 'captcha'
        await self.run_parser()
        self.assertTrue(any('проверку доступа' in text for text in self.logs))
        self.assertEqual(self.requests, [])
