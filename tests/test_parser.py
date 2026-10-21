import asyncio
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

from core.parser import TwoGisParser


class ParserTests(unittest.TestCase):
    def setUp(self):
        self.parser = TwoGisParser(lambda _: None, lambda *_: None)
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def workbook(self, filename):
        path = Path(self.tmp.name, filename)
        wb = openpyxl.Workbook()
        wb.active.append(['ID', 'Название'])
        wb.active.append(['70000001012345678', 'Кафе'])
        wb.save(path)
        wb.close()
        return str(path)

    def test_stopped_run_does_not_launch_browser(self):
        self.parser.is_running = False
        with patch('core.parser.async_playwright', side_effect=AssertionError('Browser launched')):
            self.assertEqual(asyncio.run(self.parser.run('Город', 'Кафе', 'city', '', self.tmp.name, '')), '')

    def test_stopped_backup_repeats_incomplete_page(self):
        page = self.parser._load_backup(self.workbook('STOPPED_page_5_city_query.xlsx'))
        self.assertEqual(page, 5)
        self.assertIn('70000001012345678', self.parser.results)

    def test_completed_backup_advances_page(self):
        self.assertEqual(self.parser._load_backup(self.workbook('BACKUP_page_5_city_query.xlsx')), 6)

    def test_invalid_backup_is_not_silently_ignored(self):
        with self.assertRaises(FileNotFoundError):
            self.parser._load_backup(str(Path(self.tmp.name, 'missing.xlsx')))

    def test_export_preserves_long_id_and_literal_formula(self):
        self.parser.results['70000001012345678'] = {'ID': '70000001012345678', 'Название': '=1+1'}
        path = self.parser._export_to_excel('Город', 'Кафе', self.tmp.name)
        wb = openpyxl.load_workbook(path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.active['B2'].data_type, 's')
        self.assertEqual(wb.active['B2'].value, '=1+1')
        restored = TwoGisParser(lambda _: None, lambda *_: None)
        restored._load_backup(path)
        self.assertIn('70000001012345678', restored.results)

    def test_export_sanitizes_filename(self):
        self.parser.results['1'] = {'ID': '1', 'Название': 'Кафе'}
        path = self.parser._export_to_excel('Город/Район', 'Кафе:*?', self.tmp.name)
        self.assertTrue(path)
        self.assertEqual(Path(path).parent.resolve(), Path(self.tmp.name).resolve())
        self.assertFalse(any(c in Path(path).name for c in ':*?'))

    def test_failed_export_keeps_existing_file(self):
        self.parser.results['1'] = {'ID': '1', 'Название': 'Старое'}
        path = self.parser._export_to_excel('Город', 'Кафе', self.tmp.name)
        before = Path(path).read_bytes()
        self.parser.results['1']['Название'] = 'Новое'
        with patch('openpyxl.workbook.workbook.Workbook.save', side_effect=OSError('disk error')):
            self.assertEqual(self.parser._export_to_excel('Город', 'Кафе', self.tmp.name), '')
        self.assertEqual(Path(path).read_bytes(), before)
        self.assertEqual(len(os.listdir(self.tmp.name)), 1)

    def test_numeric_response_id_updates_loaded_row(self):
        self.parser.results['123'] = {'ID': '123', 'Название': 'Кафе'}

        class Response:
            status = 200
            url = 'https://catalog.api.2gis.ru/3.0/items/byid?id=123'

            async def json(self):
                return {'result': {'items': [{'id': 123, 'name': 'Кафе',
                    'contact_groups': [{'contacts': [{'type': 'phone', 'value': '+79990000000'}]}]}]}}

        asyncio.run(self.parser._intercept_api_response(Response()))
        self.assertEqual(len(self.parser.results), 1)
        self.assertEqual(self.parser.results['123']['Телефоны'], '+79990000000')

    def test_response_suffix_does_not_duplicate_organization(self):
        class Response:
            status = 200
            url = 'https://catalog.api.2gis.ru/3.0/items/byid?id=123'

            def __init__(self, org_id):
                self.org_id = org_id

            async def json(self):
                return {'result': {'items': [{'id': self.org_id, 'name': 'Кафе'}]}}

        asyncio.run(self.parser._intercept_api_response(Response('123_old')))
        asyncio.run(self.parser._intercept_api_response(Response('123_new')))
        self.assertEqual(list(self.parser.results), ['123'])
        self.assertEqual(self.parser.results['123']['ID'], '123')

    def test_legacy_backup_normalizes_session_suffixes(self):
        path = self.workbook('STOPPED_page_1_city_query.xlsx')
        wb = openpyxl.load_workbook(path)
        wb.active['A2'] = '70000001012345678_old'
        wb.active.append(['70000001012345678_new', 'Кафе'])
        wb.save(path)
        wb.close()
        self.parser._load_backup(path)
        self.assertEqual(list(self.parser.results), ['70000001012345678'])
        self.assertEqual(self.parser.results['70000001012345678']['ID'], '70000001012345678')
        self.assertEqual(self.parser.processed_firm_ids, {'70000001012345678'})

    def test_pending_response_finishes_before_export(self):
        class Response:
            status = 200
            url = 'https://catalog.api.2gis.ru/3.0/items/byid?id=123'

            async def json(self):
                await asyncio.sleep(0.01)
                return {'result': {'items': [{'id': '123_session', 'name': 'Кафе'}]}}

        async def finish():
            self.parser._schedule_api_response(Response())
            await self.parser._wait_for_responses()
            return self.parser._export_to_excel('Город', 'Кафе', self.tmp.name)

        path = asyncio.run(finish())
        wb = openpyxl.load_workbook(path)
        self.addCleanup(wb.close)
        self.assertEqual(wb.active['A2'].value, '123')
        self.assertEqual(wb.active['B2'].value, 'Кафе')

    def test_stalled_response_does_not_block_saving(self):
        class Response:
            status = 200
            url = 'https://catalog.api.2gis.ru/3.0/items/byid?id=123'

            async def json(self):
                await asyncio.Event().wait()

        real_wait = asyncio.wait

        async def short_wait(tasks, timeout):
            return await real_wait(tasks, timeout=0.01)

        async def finish():
            self.parser._schedule_api_response(Response())
            with patch('core.parser.asyncio.wait', side_effect=short_wait):
                await asyncio.wait_for(self.parser._wait_for_responses(), timeout=0.2)
            self.assertFalse(self.parser._pending_responses)

        asyncio.run(finish())

    def test_phone_without_value_does_not_drop_organization(self):
        self.assertEqual(self.parser._extract_phone({'contact_groups': [{'contacts': [{'type': 'phone'}]}]}), 'Нет телефона')


if __name__ == '__main__':
    unittest.main()
